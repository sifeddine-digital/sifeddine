import datetime as dt
import io
import re

import openpyxl
import pytest
from werkzeug.security import generate_password_hash

from app import create_app, scoring
from app.db import execute, query
from app.referentiel import TOUS_LES_POINTS


@pytest.fixture
def app(tmp_path):
    from app import auth
    auth._failed.clear()
    app = create_app({"DATABASE": str(tmp_path / "t.sqlite3"), "SECRET_KEY": "test"})
    with app.app_context():
        for email, role in (("admin@x.ma", "admin"), ("sup@x.ma", "superviseur"), ("autre@x.ma", "superviseur")):
            execute("INSERT INTO users(email, name, role, password_hash) VALUES (?, ?, ?, ?)",
                    (email, email.split("@")[0], role, generate_password_hash("Secret123")))
        sup = query("SELECT id FROM users WHERE email = 'sup@x.ma'", one=True)["id"]
        execute("INSERT INTO stations(ship_to, name, secteur, territoire, supervisor_id) "
                "VALUES ('1001', 'SHELL A', '01 - Casa Nord', 'REN/5', ?)", (sup,))
        execute("INSERT INTO stations(ship_to, name, secteur, territoire) VALUES ('1002', 'SHELL B', '02 - Casa Sud', 'REC/3')")
    return app


def csrf_of(client, path="/login"):
    return re.search(r'name="csrf" value="([^"]+)"', client.get(path).get_data(as_text=True))[1]


def login(client, email, password="Secret123"):
    token = csrf_of(client)
    client.post("/login", data={"email": email, "password": password, "csrf": token})
    return token


def full_answers(**over):
    data = {f"r_{c}": "Oui" for c in TOUS_LES_POINTS}
    data.update({f"r_{k}": v for k, v in over.items()})
    return data


def send(client, token, station_id, week, **over):
    form = full_answers(**over) | {"csrf": token, "action": "envoyer", "visit_date": week.isoformat()}
    form.update({k: v for k, v in over.items() if not k.isupper() and not k.startswith("elim")})
    return client.post(f"/checklist/{station_id}/{week}", data=form)


# ------------------------------------------------------------------ scoring
def test_scoring_matches_sheet_rules():
    ans = {c: "Oui" for c in TOUS_LES_POINTS}
    assert scoring.score_checklist(ans, .7, .85)["statut"] == "Conforme"
    ans["elim2"] = "Non"
    assert scoring.score_checklist(ans, .7, .85)["statut"] == "Critique"
    ans = {c: "Oui" for c in TOUS_LES_POINTS}
    ans.update(H1="Non", H2="N/A")  # bloc HSSE à 0 %
    sc = scoring.score_checklist(ans, .7, .85)
    assert sc["blocs"]["hsse"] == 0 and sc["statut"] == "À corriger"
    assert sc["global"] == pytest.approx(15 / 16)


# ----------------------------------------------------------------- sécurité
def test_login_required_and_csrf(app):
    c = app.test_client()
    assert c.get("/semaine").status_code == 302
    assert c.post("/login", data={"email": "sup@x.ma", "password": "Secret123"}).status_code == 400


def test_security_headers(app):
    r = app.test_client().get("/login")
    csp = r.headers["Content-Security-Policy"]
    assert "script-src 'self'" in csp and "frame-ancestors 'none'" in csp
    assert r.headers["X-Frame-Options"] == "DENY" and r.headers["Cache-Control"] == "no-store"
    assert "onerror=" not in r.get_data(as_text=True)


def test_login_rate_limited(app):
    c = app.test_client()
    token = csrf_of(c)
    for _ in range(8):
        c.post("/login", data={"email": "sup@x.ma", "password": "faux", "csrf": token})
    r = c.post("/login", data={"email": "sup@x.ma", "password": "Secret123", "csrf": token})
    assert r.status_code == 429


def test_new_account_must_change_password(app):
    admin = app.test_client()
    token = login(admin, "admin@x.ma")
    r = admin.post("/admin/users", data={"csrf": token, "name": "Karim", "email": "karim@x.ma", "password": "Provisoire1"})
    assert r.status_code == 302
    sup = app.test_client()
    t2 = login(sup, "karim@x.ma", "Provisoire1")
    assert sup.get("/semaine").headers["Location"].endswith("/compte")
    r = sup.post("/compte", data={"csrf": t2, "old": "Provisoire1", "new": "MonSecret99", "confirm": "MonSecret99"})
    assert r.status_code == 302
    assert sup.get("/semaine").status_code == 200


def test_password_reset_closes_other_sessions(app):
    sup = app.test_client()
    login(sup, "sup@x.ma")
    assert sup.get("/semaine").status_code == 200
    admin = app.test_client()
    token = login(admin, "admin@x.ma")
    with app.app_context():
        uid = query("SELECT id FROM users WHERE email = 'sup@x.ma'", one=True)["id"]
    admin.post(f"/admin/users/{uid}", data={"csrf": token, "op": "password", "password": "Nouveau123"})
    assert sup.get("/semaine").status_code == 302  # session ancienne invalidée


def test_supervisor_sees_only_assigned_and_cannot_reach_admin(app):
    c = app.test_client()
    login(c, "sup@x.ma")
    page = c.get("/semaine").get_data(as_text=True)
    assert "SHELL A" in page and "SHELL B" not in page
    assert c.get(f"/checklist/2/{scoring.week_start()}").status_code == 403
    for path in ("/admin/", "/admin/export", "/admin/export/telecharger", "/admin/sauvegarde", "/admin/journal"):
        assert c.get(path).status_code == 403, path


# ------------------------------------------------------------- checklist
def test_send_checklist_requires_constat_when_non(app):
    c = app.test_client()
    token = login(c, "sup@x.ma")
    week = scoring.week_start()
    r = send(c, token, 1, week, P1="Non")
    assert "constat obligatoire" in r.get_data(as_text=True)
    r = send(c, token, 1, week, P1="Non", constat_proprete="Baie sale")
    assert r.status_code == 302
    with app.app_context():
        row = query("SELECT * FROM checklists WHERE station_id = 1", one=True)
        assert row["status"] == "envoye" and row["statut"] == "À corriger"  # Propreté 2/3 < 70 %
        assert query("SELECT 1 FROM journal WHERE action = 'checklist envoyée'", one=True)


def test_old_week_is_read_only_for_supervisor(app):
    c = app.test_client()
    token = login(c, "sup@x.ma")
    old = scoring.week_start() - dt.timedelta(days=21)
    r = c.post(f"/checklist/1/{old}", data=full_answers() | {"csrf": token, "action": "brouillon"})
    assert r.status_code == 403


def test_weekly_reminder_and_late_week(app):
    c = app.test_client()
    token = login(c, "sup@x.ma")
    page = c.get("/semaine").get_data(as_text=True)
    assert "1 checklist à faire" in page
    assert "Semaine dernière : 1 checklist non envoyée" in page
    assert 'class="badge">2<' in page  # semaine en cours + semaine précédente
    send(c, token, 1, scoring.week_start() - dt.timedelta(days=7))  # rattrapage S-1
    page = c.get("/semaine").get_data(as_text=True)
    assert "Semaine dernière" not in page and 'class="badge">1<' in page
    send(c, token, 1, scoring.week_start())
    page = c.get("/semaine").get_data(as_text=True)
    assert "Semaine terminée" in page and 'class="badge"' not in page


# ----------------------------------------------------------------- admin
def test_admin_assigns_stations_and_dashboard(app):
    c = app.test_client()
    token = login(c, "admin@x.ma")
    with app.app_context():
        autre = query("SELECT id FROM users WHERE email = 'autre@x.ma'", one=True)["id"]
        execute("UPDATE users SET phone = '0612345678' WHERE id = ?", (autre,))
    r = c.post(f"/admin/users/{autre}", data={"csrf": token, "op": "stations", "ids": ["1", "2"]})
    assert r.status_code == 302
    with app.app_context():
        assert {s["supervisor_id"] for s in query("SELECT supervisor_id FROM stations")} == {autre}
    page = c.get("/admin/").get_data(as_text=True)
    assert "sr-strip" in page and "Avancement de la semaine" in page
    assert "https://wa.me/212612345678?text=" in page  # relance WhatsApp
    assert c.get("/admin/journal").status_code == 200


def export(c, **params):
    r = c.get("/admin/export/telecharger", query_string=params)
    assert r.status_code == 200, r.status_code
    return openpyxl.load_workbook(io.BytesIO(r.data))


def test_export_modes(app):
    c = app.test_client()
    token = login(c, "admin@x.ma")
    w0 = scoring.week_start()
    w2 = w0 - dt.timedelta(days=14)
    send(c, token, 1, w2, elim1="Non")          # station A : ancienne checklist
    send(c, token, 2, w0, constat_service="=HYPERLINK(\"http://x\",\"clic\")", S1="Non")
    wb = export(c, mode="dernier")
    assert wb.sheetnames == ["Dashboard", "Checklist", "Invt_Eqmt", "Invt_Stck", "Historique", "Plan d'actions"]
    ck = wb["Checklist"]
    assert ck["D2"].value == "Station" and ck["AW2"].value == "Statut global"
    assert ck["F3"].value == "Non" and ck["AW3"].value.startswith("=IF(COUNTIF($F3:$H3")  # A : dernier état = S-2
    assert ck["O4"].value == "Non"                                                     # B : semaine en cours
    # un constat qui commence par « = » reste du texte (pas de formule injectée)
    cell = next(c_ for row in ck.iter_rows() for c_ in row if c_.value and "HYPERLINK" in str(c_.value))
    assert cell.data_type == "s"
    wb = export(c, mode="semaine", week=w0.isoformat())
    assert wb["Checklist"]["F3"].value is None and wb["Checklist"]["O4"].value == "Non"
    wb = export(c, mode="periode", **{"from": w2.isoformat(), "to": (w0 - dt.timedelta(days=7)).isoformat()})
    assert wb["Checklist"]["F3"].value == "Non" and wb["Checklist"]["O4"].value is None
    wb = export(c, mode="dernier", secteurs="02 - Casa Sud")
    assert wb["Checklist"]["D3"].value == "SHELL B" and wb["Checklist"]["D4"].value is None


def test_backup_download(app):
    c = app.test_client()
    login(c, "admin@x.ma")
    r = c.get("/admin/sauvegarde")
    assert r.status_code == 200 and r.data[:15] == b"SQLite format 3"


def test_export_import_round_trip(app, tmp_path):
    c = app.test_client()
    token = login(c, "admin@x.ma")
    week = scoring.week_start()
    send(c, token, 1, week, S2="Non", constat_service="Pas de conseil")
    data = c.get("/admin/export/telecharger", query_string={"mode": "semaine", "week": week.isoformat()}).data
    app2 = create_app({"DATABASE": str(tmp_path / "t2.sqlite3"), "SECRET_KEY": "test"})
    with app2.app_context():
        from app.importer import import_fiche
        report = import_fiche(io.BytesIO(data))
        assert "2 lues" in report[0]
        row = query("SELECT * FROM checklists", one=True)
        assert row["week_start"] == week.isoformat() and row["statut"] == "À corriger"


def test_switch_account_from_login_page(app):
    c = app.test_client()
    login(c, "admin@x.ma")
    r = c.get("/login")
    assert r.status_code == 200 and "Vous êtes connecté en tant que" in r.get_data(as_text=True)
    login(c, "sup@x.ma")  # même navigateur : on passe au compte superviseur
    page = c.get("/semaine").get_data(as_text=True)
    assert "SHELL A" in page and "SHELL B" not in page
    assert c.get("/admin/").status_code == 403


def test_station_fiche_excel(app):
    c = app.test_client()
    token = login(c, "admin@x.ma")
    send(c, token, 1, scoring.week_start())
    wb = export_station = openpyxl.load_workbook(io.BytesIO(c.get("/admin/stations/1/fiche.xlsx").data))
    ck = wb["Checklist"]
    assert ck["D3"].value == "SHELL A" and ck["D4"].value is None and ck["F3"].value == "Oui"
    assert export_station.sheetnames[0] == "Dashboard"
