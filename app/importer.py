"""Import de la fiche Excel « LOS Inventaire & Checklist » (repère les colonnes par leur en-tête)."""
import datetime as dt
import json
import re
import unicodedata
from collections import Counter

import openpyxl

from . import scoring
from .db import get_db, seuils
from .referentiel import (BLOCS, EQUIPEMENTS, STATUTS_RELEVE, STOCK_FREQUENCES,
                          STOCK_METHODES, TOUS_LES_POINTS)


def _norm(s):
    s = unicodedata.normalize("NFKD", str(s or "")).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]+", " ", s).strip()


def _header(ws, must="code ship to", max_row=10):
    for r in range(1, max_row + 1):
        cols = {_norm(ws.cell(r, c).value): c for c in range(1, ws.max_column + 1) if ws.cell(r, c).value}
        if must in cols:
            return r, cols
    raise ValueError(f"en-tête « Code Ship-to » introuvable dans l'onglet {ws.title}")


def _yn(v, allow_na=False):
    v = str(v or "").strip().lower()
    if v == "oui":
        return "Oui"
    if v == "non":
        return "Non"
    if allow_na and v in ("n/a", "na"):
        return "N/A"
    return ""


def _ship(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def _date(v):
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    m = re.match(r"^\s*(\d{1,2})[./-](\d{1,2})[./-](\d{4})\s*$", str(v or ""))
    if m:
        try:
            return dt.date(int(m[3]), int(m[2]), int(m[1]))
        except ValueError:
            return None
    return None


def _text(v):
    v = "" if v is None else str(v).strip()
    return "" if v in ("-", "—") or v.startswith("=") else v


def _month_of(ws):
    for r in range(1, 4):
        for c in range(1, 4):
            d = _date(ws.cell(r, c).value)
            if d:
                return d.replace(day=1)
    return scoring.month_start()


def import_fiche(stream, with_data=True, user_id=None):
    wb = openpyxl.load_workbook(stream)
    db = get_db()
    report = []
    names = {_norm(n): n for n in wb.sheetnames}

    # ---- stations (ordre de l'inventaire, puis checklist, puis stock)
    seen, order = {}, 0
    for key in ("invt eqmt", "checklist", "invt stck"):
        if key not in names:
            continue
        ws = wb[names[key]]
        hr, cols = _header(ws)
        for r in range(hr + 1, ws.max_row + 1):
            ship = _ship(ws.cell(r, cols["code ship to"]).value)
            if not ship.isdigit() or ship in seen:
                continue
            order += 1
            seen[ship] = (_text(ws.cell(r, cols["secteur"]).value), _text(ws.cell(r, cols.get("territoire", 2)).value),
                          _text(ws.cell(r, cols["station"]).value), order)
    created = 0
    for ship, (secteur, territoire, name, order) in seen.items():
        cur = db.execute("UPDATE stations SET secteur=?, territoire=?, name=?, sort_order=? WHERE ship_to=?",
                         (secteur, territoire, name, order, ship))
        if cur.rowcount == 0:
            db.execute("INSERT INTO stations(ship_to, name, secteur, territoire, sort_order) VALUES (?,?,?,?,?)",
                       (ship, name, secteur, territoire, order))
            created += 1
    db.commit()
    report.append(f"Stations : {len(seen)} lues, {created} nouvelles, {len(seen) - created} mises à jour.")
    if not with_data:
        return report
    ids = {r["ship_to"]: r["id"] for r in db.execute("SELECT id, ship_to FROM stations")}

    # ---- inventaire équipements
    if "invt eqmt" in names:
        ws = wb[names["invt eqmt"]]
        month = _month_of(ws)
        hr, cols = _header(ws)
        eq_cols = {key: cols.get(_norm(label)) for _, _, items in EQUIPEMENTS for key, label, _ in items}
        missing = [k for k, c in eq_cols.items() if not c]
        n = 0
        for r in range(hr + 1, ws.max_row + 1):
            sid = ids.get(_ship(ws.cell(r, cols["code ship to"]).value))
            if not sid:
                continue
            statut = _text(ws.cell(r, cols.get("statut releve", 0) or 5).value)
            baie, shsc = _yn(ws.cell(r, cols["baie"]).value), _yn(ws.cell(r, cols["shsc"]).value)
            items = {k: _yn(ws.cell(r, c).value) for k, c in eq_cols.items() if c and _yn(ws.cell(r, c).value)}
            if not (statut or baie or items):
                continue
            db.execute("""INSERT INTO inventaires(station_id, month, statut, baie, shsc, items, updated_by)
                          VALUES (?,?,?,?,?,?,?) ON CONFLICT(station_id, month) DO UPDATE SET statut=excluded.statut,
                          baie=excluded.baie, shsc=excluded.shsc, items=excluded.items, updated_at=datetime('now')""",
                       (sid, month.isoformat(), statut if statut in STATUTS_RELEVE else "", baie, shsc,
                        json.dumps(items), user_id))
            n += 1
        db.commit()
        report.append(f"Inventaire équipements ({scoring.month_label(month)}) : {n} station(s)."
                      + (f" Colonnes non reconnues : {', '.join(missing)}." if missing else ""))

    # ---- stock
    if "invt stck" in names:
        ws = wb[names["invt stck"]]
        month = _month_of(ws)
        hr, cols = _header(ws)

        def col(*prefixes):
            for name, c in cols.items():
                if any(name.startswith(p) for p in prefixes):
                    return c
            return None
        c_fiable, c_meth, c_date = col("stock fiable"), col("methode"), col("date du dernier")
        c_freq, c_sys, c_phy, c_com = col("frequence"), col("stock systeme"), col("stock physique"), col("commentaire")
        n = 0
        for r in range(hr + 1, ws.max_row + 1):
            sid = ids.get(_ship(ws.cell(r, cols["code ship to"]).value))
            get = lambda c: ws.cell(r, c).value if c else None  # noqa: E731
            if not sid or not any(get(c) not in (None, "") for c in (c_fiable, c_meth, c_date, c_freq, c_sys, c_com)):
                continue
            d = _date(get(c_date))
            meth, freq = _text(get(c_meth)), _text(get(c_freq))
            num = lambda v: float(v) if isinstance(v, (int, float)) else None  # noqa: E731
            db.execute("""INSERT INTO stocks(station_id, month, fiable, methode, date_inventaire, frequence,
                          stock_physique, stock_systeme, commentaire, updated_by) VALUES (?,?,?,?,?,?,?,?,?,?)
                          ON CONFLICT(station_id, month) DO UPDATE SET fiable=excluded.fiable, methode=excluded.methode,
                          date_inventaire=excluded.date_inventaire, frequence=excluded.frequence,
                          stock_physique=excluded.stock_physique, stock_systeme=excluded.stock_systeme,
                          commentaire=excluded.commentaire, updated_at=datetime('now')""",
                       (sid, month.isoformat(), _yn(get(c_fiable)), meth if meth in STOCK_METHODES else "",
                        d.isoformat() if d else "", freq if freq in STOCK_FREQUENCES else "",
                        num(get(c_phy)), num(get(c_sys)), _text(get(c_com)), user_id))
            n += 1
        db.commit()
        report.append(f"Stock ({scoring.month_label(month)}) : {n} station(s).")

    # ---- checklist
    if "checklist" in names:
        ws = wb[names["checklist"]]
        hr, cols = _header(ws)
        code_cols = {}
        for code in TOUS_LES_POINTS:
            prefix = _norm(code if not code.startswith("elim") else f"ELIM {code[-1]}")
            code_cols[code] = next((c for n_, c in cols.items() if n_.startswith(prefix + " ")), None)
        bloc_cols = {}
        for cle, titre, _ in BLOCS:
            nom = _norm(titre.split(". ", 1)[1])
            bloc_cols[cle] = (cols.get(f"constat {nom}"), cols.get(f"action {nom}"))
        c_date, c_resp, c_com = cols.get("date de visite"), cols.get("responsable echeance"), cols.get("commentaire")
        parsed = []
        for r in range(hr + 1, ws.max_row + 1):
            sid = ids.get(_ship(ws.cell(r, cols["code ship to"]).value))
            if not sid:
                continue
            answers = {code: _yn(ws.cell(r, c).value, allow_na=True) for code, c in code_cols.items() if c}
            answers = {k: v for k, v in answers.items() if v}
            if not answers:
                continue
            data = {
                "answers": answers,
                "constats": {k: _text(ws.cell(r, c[0]).value) if c[0] else "" for k, c in bloc_cols.items()},
                "actions": {k: _text(ws.cell(r, c[1]).value) if c[1] else "" for k, c in bloc_cols.items()},
                "responsable": _text(ws.cell(r, c_resp).value) if c_resp else "",
                "commentaire": _text(ws.cell(r, c_com).value) if c_com else "",
                "visit_date": "",
            }
            d = _date(ws.cell(r, c_date).value) if c_date else None
            if d:
                data["visit_date"] = d.isoformat()
            parsed.append((sid, d, data))
        dated = Counter(scoring.week_start(d) for _, d, _ in parsed if d)
        fallback = dated.most_common(1)[0][0] if dated else scoring.week_start(
            _month_of(wb[names["invt eqmt"]]) if "invt eqmt" in names else None)
        seuil_bloc, seuil_global = seuils()
        undated = 0
        for sid, d, data in parsed:
            week = scoring.week_start(d) if d else fallback
            undated += 0 if d else 1
            sc = scoring.score_checklist(data["answers"], seuil_bloc, seuil_global)
            complete = all(c in data["answers"] for c in TOUS_LES_POINTS)
            db.execute("""INSERT INTO checklists(station_id, week_start, visit_date, data, status, score_global, statut,
                          updated_by, submitted_at) VALUES (?,?,?,?,?,?,?,?, CASE WHEN ? THEN datetime('now') END)
                          ON CONFLICT(station_id, week_start) DO UPDATE SET visit_date=excluded.visit_date,
                          data=excluded.data, status=excluded.status, score_global=excluded.score_global,
                          statut=excluded.statut, updated_at=datetime('now')""",
                       (sid, week.isoformat(), data["visit_date"] or None, json.dumps(data, ensure_ascii=False),
                        "envoye" if complete else "brouillon", sc["global"], sc["statut"], user_id, complete))
        db.commit()
        msg = f"Checklist : {len(parsed)} station(s) avec réponses."
        if undated:
            msg += f" {undated} sans date de visite → rangées en {scoring.week_label(fallback)}."
        report.append(msg)
    return report
